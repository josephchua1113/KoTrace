#!/usr/bin/env python3
"""Compare an analyzer workbook against a sample's expected_*.csv files.

    python3 check_expected.py --xlsx out.xlsx --sample samples/shop-app
"""
import argparse
import csv
import os
import sys

from openpyxl import load_workbook

# files that exist in the sample but must never appear anywhere in the report
FORBIDDEN = ["ProductRepoHelper", "Legacyservice",
             "FakeProfileService", "MockCartService", "GeneratedApiService"]


def sheet_rows(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    return rows[0], [tuple("" if v is None else str(v) for v in r) for r in rows[1:]]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sample", required=True)
    a = ap.parse_args()
    wb = load_workbook(a.xlsx)
    fails = []

    # 1. matrix: exact rows, exact order
    _, got = sheet_rows(wb, "Traceability Matrix")
    with open(os.path.join(a.sample, "expected_matrix.csv"), newline="") as f:
        exp = [tuple(r) for r in list(csv.reader(f))[1:]]
    if got != exp:
        for r in exp:
            if r not in got:
                fails.append("missing matrix row: %s" % (r,))
        for r in got:
            if r not in exp:
                fails.append("unexpected matrix row: %s" % (r,))
        if sorted(got) == sorted(exp):
            fails.append("matrix rows correct but in the wrong order")

    # 1b. broken paths: proven cells + where each chain stopped (first 15 columns)
    bp = os.path.join(a.sample, "expected_broken.csv")
    if os.path.exists(bp):
        _, got_b = sheet_rows(wb, "Broken Paths")
        got_b = [r[:15] for r in got_b]
        with open(bp, newline="") as f:
            exp_b = [tuple(r) for r in list(csv.reader(f))[1:]]
        for r in exp_b:
            if r not in got_b:
                fails.append("missing broken path: %s" % (r,))
        for r in got_b:
            if r not in exp_b:
                fails.append("unexpected broken path: %s" % (r,))
        if sorted(got_b) == sorted(exp_b) and got_b != exp_b:
            fails.append("broken paths correct but in the wrong order")

    # 2. Step 0 filter
    _, files = sheet_rows(wb, "Filtered Files")
    got_files = sorted((r[0], r[1]) for r in files)
    with open(os.path.join(a.sample, "expected_files.csv"), newline="") as f:
        exp_files = sorted(tuple(r) for r in list(csv.reader(f))[1:])
    for r in sorted(set(exp_files) - set(got_files)):
        fails.append("file missing from filter: %s" % (r,))
    for r in sorted(set(got_files) - set(exp_files)):
        fails.append("file wrongly accepted by filter: %s" % (r,))

    # 3. Unresolved entries
    _, unres = sheet_rows(wb, "Unresolved")
    with open(os.path.join(a.sample, "expected_unresolved.csv"), newline="") as f:
        for typ, needle in list(csv.reader(f))[1:]:
            if not any(r[1] == typ and (needle in r[0] or needle in r[4]) for r in unres):
                fails.append("expected Unresolved [%s] mentioning '%s'" % (typ, needle))

    # 4. no excluded file anywhere in the workbook
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                for bad in FORBIDDEN:
                    if v is not None and bad in str(v):
                        fails.append("excluded file '%s' appears in sheet %s" % (bad, ws.title))

    if fails:
        print("FAIL (%d)" % len(fails))
        for x in fails:
            print("  - " + x)
        return 1
    print("PASS: %d matrix rows, %s broken paths, %d filtered files, %d expected Unresolved entries found"
          % (len(exp), len(exp_b) if os.path.exists(bp) else "no", len(exp_files), sum(1 for _ in open(os.path.join(a.sample, "expected_unresolved.csv"))) - 1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
